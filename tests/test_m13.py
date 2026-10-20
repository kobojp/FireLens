from __future__ import annotations

import io
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app.db import connect
from backend.app.main import create_app
from backend.app.reconcile import (
    build_reconciliation,
    export_reconciliation,
    import_ledger,
    parse_excel_bytes,
    parse_text_rows,
    preview_excel_bytes,
    preview_text_rows,
)
from backend.app.scanner import scan_root


def _workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _add_root(conn, source: Path) -> int:
    with conn:
        cursor = conn.execute(
            "INSERT INTO root(label, path, is_cloud_stream) VALUES ('M13', ?, 0)",
            (str(source.resolve()),),
        )
    root_id = int(cursor.lastrowid)
    scan_root(conn, root_id)
    return root_id


def _item(source: Path, code: str, files: tuple[str, ...], building: str = "中正樓") -> None:
    path = source / "滅火器" / "2026年" / "10月份更換" / building / code
    path.mkdir(parents=True, exist_ok=True)
    for filename in files:
        (path / filename).write_bytes(filename.encode("utf-8"))


def test_text_parser_accepts_device_only_and_two_columns() -> None:
    single = parse_text_rows("動-01-15\n動-03-08\n致-B1-67\n中-B2-05")
    assert [row["device_no"] for row in single] == [
        "動-01-15",
        "動-03-08",
        "致-B1-67",
        "中-B2-05",
    ]
    assert all(row["section"] == "" for row in single)

    double = parse_text_rows("致德樓8F\t致-08-19\n中正樓B2F\t中-B2-05")
    assert double[0]["section"] == "致德樓8F"
    assert double[0]["device_no"] == "致-08-19"
    assert double[1]["section"] == "中正樓B2F"


def test_text_preview_normalizes_and_reports_duplicates_invalid_and_warning() -> None:
    preview = preview_text_rows(
        "棟別樓層\t設備編號\n致德樓8F\t 致 - 08 - 19 \n錯誤樓8F\t致-08-19\n中正樓B2F\tBADVALUE\n"
    )
    assert preview["has_header"] is True
    assert preview["total_rows"] == 3
    assert preview["valid_count"] == 2
    assert preview["duplicate_count"] == 2
    assert preview["invalid_count"] == 1
    rows = preview["rows"]
    assert rows[0]["device_no"] == "致-08-19"
    assert rows[0]["duplicate"] is True
    assert any("重複" in warning for warning in rows[0]["warnings"])
    assert any("不一致" in warning for warning in rows[1]["warnings"])
    assert rows[2]["valid"] is False
    assert rows[2]["error"]


def test_excel_accepts_headers_without_headers_one_two_columns_and_extra_columns() -> None:
    with_header = parse_excel_bytes(
        _workbook(
            [
                ["棟別樓層", "設備編號", "備註"],
                ["致德樓8F", "致-08-19", "忽略我"],
                ["中正樓B2F", "中-B2-05", "其他"],
            ]
        )
    )
    assert [(row["section"], row["device_no"]) for row in with_header] == [
        ("致德樓8F", "致-08-19"),
        ("中正樓B2F", "中-B2-05"),
    ]

    no_header = parse_excel_bytes(_workbook([["致德樓B1F", "致-B1-67"], ["中正樓B2F", "中-B2-28"]]))
    assert no_header[0]["device_no"] == "致-B1-67"
    assert no_header[1]["section"] == "中正樓B2F"

    one_column = parse_excel_bytes(_workbook([["致-08-69"], ["思-06-15"]]))
    assert [row["device_no"] for row in one_column] == ["致-08-69", "思-06-15"]
    assert all(row["section"] == "" for row in one_column)


def test_old_excel_headers_remain_compatible_and_legacy_fields_survive(db_path: Path) -> None:
    rows = parse_excel_bytes(
        _workbook(
            [
                ["區段", "設備名稱", "型式規格", "上次更換日"],
                ["中正樓B2F", "中-B2-05", "10P", "2026/10/01"],
            ]
        )
    )
    assert rows[0]["section"] == "中正樓B2F"
    assert rows[0]["specification"] == "10P"
    assert rows[0]["last_replacement_date"] == "2026/10/01"
    conn = connect(db_path)
    try:
        ledger_id = import_ledger(conn, name="legacy", source_type="text", rows=rows)
        stored = conn.execute(
            "SELECT section, device_no, specification, last_replacement_date, raw_json "
            "FROM ledger_row WHERE ledger_id=?",
            (ledger_id,),
        ).fetchone()
        assert stored is not None
        assert stored["specification"] == "10P"
        assert stored["last_replacement_date"] == "2026/10/01"
        assert '"specification": "10P"' in stored["raw_json"]
    finally:
        conn.close()


def test_preview_api_and_text_import_keep_original_text(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("FIRELENS_DATA_DIR", str(tmp_path / "appdata"))
    text = "致德樓8F\t致-08-19\n中正樓B2F\t中-B2-05\n壞資料"
    with TestClient(create_app()) as client:
        preview = client.post(
            "/api/ledgers/text/preview",
            json={"name": "2026/10 滅火器清冊", "text": text},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid_count"] == 2
        assert preview.json()["invalid_count"] == 1

        created = client.post(
            "/api/ledgers/text",
            json={"name": "2026/10 滅火器清冊", "text": text},
        )
        assert created.status_code == 201, created.text
        detail = client.get(f"/api/ledgers/{created.json()['id']}").json()
        assert detail["source_text"] == text
        assert detail["row_count"] == 2


def test_excel_preview_api_supports_user_two_column_format(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("FIRELENS_DATA_DIR", str(tmp_path / "appdata"))
    data = _workbook(
        [
            ["致德樓8F", "致-08-19"],
            ["致德樓8F", "致-08-69"],
            ["致德樓B1F", "致-B1-67"],
            ["中正樓B2F", "中-B2-05"],
            ["中正樓B2F", "中-B2-28"],
        ]
    )
    with TestClient(create_app()) as client:
        response = client.post(
            "/api/ledgers/excel/preview",
            files={
                "file": (
                    "滅火器.xlsx",
                    data,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert response.status_code == 200, response.text
        payload = response.json()
        assert payload["valid_count"] == 5
        assert payload["source_filename"] == "滅火器.xlsx"
        assert payload["rows"][2]["device_no"] == "致-B1-67"


def test_reconciliation_has_broad_classifications_and_missing_export(
    db_path: Path, tmp_path: Path
) -> None:
    source = tmp_path / "photos"
    _item(source, "中-B2-05", ("編號.jpg", "藥劑.jpg", "完成.jpg"))
    _item(source, "中-B2-28", ("編號.jpg", "藥劑.jpg"))
    _item(source, "中-B2-30", ())
    _item(source, "中-B2-31", ("編號.jpg", "藥劑.jpg", "完成.jpg", "46459.jpg"))

    conn = connect(db_path)
    try:
        root_id = _add_root(conn, source)
        ledger_id = import_ledger(
            conn,
            name="M13",
            source_type="text",
            rows=parse_text_rows(
                "中正樓B2F\t中-B2-05\n"
                "中正樓B2F\t中-B2-28\n"
                "中正樓B2F\t中-B2-30\n"
                "中正樓B2F\t中-B2-31\n"
                "中正樓B2F\t中-B2-99"
            ),
        )
        result = build_reconciliation(conn, root_id=root_id, ledger_id=ledger_id)
        by_no = {row["device_no"]: row for row in result["rows"] if row["kind"] == "ledger"}
        assert by_no["中-B2-05"]["classification"] == "complete"
        assert by_no["中-B2-28"]["classification"] == "partial_missing"
        assert by_no["中-B2-30"]["classification"] == "no_photos"
        assert by_no["中-B2-31"]["classification"] == "unmapped_legacy"
        assert by_no["中-B2-99"]["classification"] == "no_folder"
        assert result["classification_counts"] == {
            "complete": 1,
            "no_folder": 1,
            "no_photos": 1,
            "partial_missing": 1,
            "unmapped_legacy": 1,
        }

        content, _, filename = export_reconciliation(result, "xlsx")
        assert filename == "補拍清單.xlsx"
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            values = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
        assert values[0][:5] == ("類別", "設備編號", "棟別樓層", "核對分類", "缺少照片／狀態")
        exported_ids = {row[1] for row in values[1:]}
        assert "中-B2-05" not in exported_ids
        assert {"中-B2-28", "中-B2-30", "中-B2-31", "中-B2-99"}.issubset(exported_ids)
    finally:
        conn.close()


def test_excel_preview_function_reports_header_and_validity() -> None:
    preview = preview_excel_bytes(
        _workbook([["棟別樓層", "設備編號"], ["致德樓8F", "致-08-19"], ["", "BAD"]])
    )
    assert preview["has_header"] is True
    assert preview["valid_count"] == 1
    assert preview["invalid_count"] == 1


def test_installer_uses_repo_traditional_chinese_language_resource() -> None:
    root = Path(__file__).parents[1]
    script = (root / "packaging" / "FireLens.iss").read_text(encoding="utf-8")
    language_path = root / "packaging" / "lang" / "ChineseTraditional.isl"
    assert language_path.is_file()
    language = language_path.read_text(encoding="utf-8-sig")
    assert 'Name: "chinesetraditional"; MessagesFile: "lang\\ChineseTraditional.isl"' in script
    assert "ShowLanguageDialog=no" in script
    for message in (
        "ButtonBack=",
        "ButtonNext=",
        "ButtonCancel=",
        "WizardSelectDir=",
        "WizardSelectTasks=",
        "WizardReady=",
        "WizardInstalling=",
        "FinishedHeadingLabel=",
        "ButtonFinish=",
        "DownloadingLabel2=",
        "ExtractingLabel=",
    ):
        assert message in language
