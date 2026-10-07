from __future__ import annotations

import io
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app.db import connect
from backend.app.main import create_app
from backend.app.reconcile import (
    build_reconciliation,
    export_reconciliation,
    import_ledger,
    normalize_building,
    parse_excel_bytes,
    parse_text_rows,
    seed_default_aliases,
)
from backend.app.scanner import scan_root


def workbook_bytes(rows: list[tuple[object, object, object, object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["區段", "設備名稱", "型式規格", "上次更換日"])
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def add_complete_item(root: Path, period: str, building: str, code: str) -> Path:
    item = root / "滅火器" / "2026年" / period / building / code
    item.mkdir(parents=True, exist_ok=True)
    (item / "編號.jpg").write_bytes(b"number")
    (item / "藥劑.jpg").write_bytes(b"agent")
    (item / "完成.jpg").write_bytes(b"done")
    return item


def add_root_and_scan(db_path: Path, source: Path) -> int:
    conn = connect(db_path)
    try:
        with conn:
            cursor = conn.execute(
                "INSERT INTO root(label, path, is_cloud_stream) VALUES ('M3', ?, 0)",
                (str(source.resolve()),),
            )
        root_id = int(cursor.lastrowid)
        scan_root(conn, root_id)
        return root_id
    finally:
        conn.close()


def test_near_number_is_conservative() -> None:
    from backend.app.reconcile import _near_candidates

    candidates = ["中-01-130", "中-01-14", "動-01-13", "中-02-13", "中-01-013"]
    assert _near_candidates("中-01-13", candidates) == ["中-01-013", "中-01-130"]
    assert _near_candidates("動-01-15", ["一-01-15", "動-01-16", "動-03-15"]) == []


def test_excel_and_text_parser_accept_special_floor_ids() -> None:
    data = workbook_bytes(
        [
            ("第一區1X", "一-1X-01", "10P", "2023-09-01"),
            ("機電區1M", "機-1M-01", "10P", "2023-09-01"),
            ("C棟RF", "C-RF-01", "10P", "2023-09-01"),
        ]
    )
    rows = parse_excel_bytes(data)
    assert [row["device_no"] for row in rows] == ["一-1X-01", "機-1M-01", "C-RF-01"]

    text = "區段\t設備名稱\t型式規格\t上次更換日\n第一區1X\t一-1X-01\t10P\t2023/09/01"
    parsed = parse_text_rows(text)
    assert parsed[0]["device_no"] == "一-1X-01"
    assert parsed[0]["section"] == "第一區1X"


def test_alias_normalization_uses_safe_default_rules(db_path: Path) -> None:
    conn = connect(db_path)
    try:
        seed_default_aliases(conn)
        aliases = {
            row[0]: row[1]
            for row in conn.execute(
                "SELECT source_value, target_value FROM alias_map WHERE alias_type='building'"
            )
        }
        assert normalize_building("第一區B1F", aliases) == "一區"
        assert normalize_building("第二區1F", aliases) == "二區"
        assert normalize_building("B棟異體8F", aliases) == "B棟"
        assert "未設定區域" not in aliases
    finally:
        conn.close()


def test_reconciliation_statuses_cross_month_aliases_and_exports(
    db_path: Path, sample_source: Path
) -> None:
    add_complete_item(sample_source, "9月份更換", "一區", "一-1X-01")
    add_complete_item(sample_source, "9月份更換", "中正樓", "中-01-130")

    dirty = sample_source / "滅火器" / "2026年" / "9月份更換" / "中正樓" / "中-02-01"
    dirty.mkdir(parents=True)
    (dirty / "編輯.jpg").write_bytes(b"number-typo")
    (dirty / "藥劑.jpg").write_bytes(b"agent")
    (dirty / "完成.jpg").write_bytes(b"done")

    zero = sample_source / "滅火器" / "2026年" / "9月份更換" / "中正樓" / "中-03-01"
    zero.mkdir(parents=True)
    (zero / "編號.jpg").write_bytes(b"")
    (zero / "藥劑.jpg").write_bytes(b"agent")
    (zero / "完成.jpg").write_bytes(b"done")

    add_complete_item(sample_source, "9月份更換", "中正樓", "中-04-01")
    add_complete_item(sample_source, "10月份更換", "中正樓", "中-04-01")
    add_complete_item(sample_source, "10月份更換", "中正樓", "額-01-01")
    add_complete_item(sample_source, "壓力錶壞", "中正樓", "特-01-01")
    regular_legacy_box = sample_source / "放置盒" / "2026年" / "9月份更換" / "中正樓" / "箱舊-01"
    regular_legacy_box.mkdir(parents=True)
    (regular_legacy_box / "新增1.jpg").write_bytes(b"legacy")
    (regular_legacy_box / "完成.jpg").write_bytes(b"legacy-done")

    root_id = add_root_and_scan(db_path, sample_source)
    rows = [
        {
            "section": "中正樓1F",
            "device_no": "中-01-01",
            "specification": "10P",
            "last_replacement_date": "2023-09-01",
        },
        {
            "section": "中正樓1F",
            "device_no": "中-01-02",
            "specification": "10P",
            "last_replacement_date": "2023-09-01",
        },
        {
            "section": "第一區1X",
            "device_no": "一-1X-01",
            "specification": "10P",
            "last_replacement_date": "2023-09-01",
        },
        {
            "section": "中正樓1F",
            "device_no": "中-01-13",
            "specification": "10P",
            "last_replacement_date": "2023-09-01",
        },
        {
            "section": "中正樓2F",
            "device_no": "中-02-01",
            "specification": "10P",
            "last_replacement_date": "2023-09-01",
        },
        {
            "section": "中正樓3F",
            "device_no": "中-03-01",
            "specification": "10P",
            "last_replacement_date": "2023-09-01",
        },
        {
            "section": "中正樓4F",
            "device_no": "中-04-01",
            "specification": "10P",
            "last_replacement_date": "2023-09-01",
        },
        {
            "section": "中正樓5F",
            "device_no": "中-99-99",
            "specification": "10P",
            "last_replacement_date": "2023-09-01",
        },
    ]

    conn = connect(db_path)
    try:
        seed_default_aliases(conn)
        ledger_id = import_ledger(conn, name="M3 synthetic", source_type="text", rows=rows)
        result = build_reconciliation(conn, root_id=root_id, ledger_id=ledger_id)
        by_no = {row["device_no"]: row for row in result["rows"] if row["kind"] == "ledger"}
        assert by_no["中-01-01"]["status"] == "complete"
        assert by_no["中-01-02"]["status"] == "empty_folder"
        assert by_no["一-1X-01"]["status"] == "complete"
        assert by_no["一-1X-01"]["building_match"] is True
        assert by_no["中-01-13"]["status"] == "near_number"
        assert by_no["中-01-13"]["near_candidates"] == ["中-01-130"]
        assert by_no["中-02-01"]["status"] == "invalid_naming"
        assert by_no["中-02-01"]["missing_slots"] == ["編號"]
        assert by_no["中-02-01"]["naming_suggestions"][0]["suggestion"] == "編號.jpg"
        assert by_no["中-03-01"]["status"] == "zero_byte"
        assert by_no["中-04-01"]["status"] == "duplicate_number"
        assert by_no["中-99-99"]["status"] == "no_folder"

        extras = [row["device_no"] for row in result["rows"] if row["status"] == "extra_folder"]
        assert "額-01-01" in extras
        special = [row for row in result["rows"] if row["status"] == "special"]
        assert any(row["device_no"] == "特-01-01" for row in special)
        slot_only = {
            (row["category"], row["device_no"]): row
            for row in result["rows"]
            if row["kind"] == "slot_only"
        }
        assert slot_only[("box", "中-02-C124")]["status"] == "special"
        assert slot_only[("box", "中-02-C124")]["missing_slots"] == []
        assert slot_only[("box", "箱舊-01")]["status"] == "legacy"
        assert slot_only[("box", "箱舊-01")]["missing_slots"] == []
        assert slot_only[("lamp", "D-05-04")]["status"] == "legacy"
        assert slot_only[("lamp", "D-05-04")]["missing_slots"] == []
        assert result["total"] == 8
        assert result["complete"] == 2
        assert result["completion_rate"] == 25.0

        restricted = build_reconciliation(
            conn, root_id=root_id, ledger_id=ledger_id, period="10月份更換"
        )
        restricted_row = next(
            row
            for row in restricted["rows"]
            if row["kind"] == "ledger" and row["device_no"] == "中-04-01"
        )
        assert restricted_row["status"] == "complete"
        assert restricted_row["found"][0]["period"] == "10月份更換"

        csv_bytes, media_type, _ = export_reconciliation(result, "csv")
        assert media_type.startswith("text/csv")
        assert "中-99-99" in csv_bytes.decode("utf-8-sig")
        xlsx_bytes, xlsx_type, _ = export_reconciliation(result, "xlsx")
        assert "spreadsheetml" in xlsx_type
        workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
        try:
            assert workbook.active.max_row > 1
        finally:
            workbook.close()
    finally:
        conn.close()


def test_cloud_placeholder_status(db_path: Path, sample_source: Path) -> None:
    root_id = add_root_and_scan(db_path, sample_source)
    conn = connect(db_path)
    try:
        seed_default_aliases(conn)
        item = conn.execute("SELECT id FROM item WHERE code='中-01-01'").fetchone()
        assert item is not None
        with conn:
            conn.execute(
                "UPDATE photo SET cloud_placeholder=1 WHERE item_id=? AND slot='編號'",
                (item[0],),
            )
        ledger_id = import_ledger(
            conn,
            name="cloud",
            source_type="text",
            rows=[
                {
                    "section": "中正樓1F",
                    "device_no": "中-01-01",
                    "specification": "10P",
                    "last_replacement_date": "",
                }
            ],
        )
        result = build_reconciliation(conn, root_id=root_id, ledger_id=ledger_id)
        assert result["rows"][0]["status"] == "cloud_not_downloaded"
    finally:
        conn.close()


def test_m3_api_excel_text_reconcile_and_download(
    monkeypatch, tmp_path: Path, sample_source: Path
) -> None:
    monkeypatch.setenv("FIRELENS_DATA_DIR", str(tmp_path / "appdata"))
    app = create_app()
    with TestClient(app) as client:
        root_response = client.post(
            "/api/roots", json={"path": str(sample_source), "label": "M3 API"}
        )
        root_id = root_response.json()["id"]
        assert client.post(f"/api/roots/{root_id}/scan").status_code == 200

        data = workbook_bytes(
            [
                ("中正樓1F", "中-01-01", "10P", "2023-09-01"),
                ("中正樓1F", "中-99-99", "10P", "2023-09-01"),
            ]
        )
        imported = client.post(
            "/api/ledgers/excel",
            files={
                "file": (
                    "ledger.xlsx",
                    data,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert imported.status_code == 201
        assert imported.json()["row_count"] == 2
        ledger_id = imported.json()["id"]

        text_import = client.post(
            "/api/ledgers/text",
            json={"name": "貼上清冊", "text": "區段\t設備名稱\n中正樓1F\t中-01-01"},
        )
        assert text_import.status_code == 201
        assert text_import.json()["row_count"] == 1

        aliases = client.get("/api/aliases")
        assert aliases.status_code == 200
        assert any(row["source_value"] == "第一區" for row in aliases.json())

        result = client.post(
            "/api/reconcile",
            json={"root_id": root_id, "ledger_id": ledger_id, "period": None},
        )
        assert result.status_code == 200
        assert result.json()["total"] == 2
        assert "中-99-99" in result.json()["reshoot_text"]

        marked = client.post(
            "/api/reconcile/reshot",
            json={
                "root_id": root_id,
                "ledger_id": ledger_id,
                "category_code": "extinguisher",
                "device_no": "中-99-99",
                "marked": True,
            },
        )
        assert marked.status_code == 200
        rerun = client.post(
            "/api/reconcile",
            json={"root_id": root_id, "ledger_id": ledger_id, "period": None},
        ).json()
        marked_row = next(row for row in rerun["rows"] if row["device_no"] == "中-99-99")
        assert marked_row["marked_reshot"] is True

        unmarked = client.post(
            "/api/reconcile/reshot",
            json={
                "root_id": root_id,
                "ledger_id": ledger_id,
                "category_code": "extinguisher",
                "device_no": "中-99-99",
                "marked": False,
            },
        )
        assert unmarked.status_code == 200
        rerun = client.post(
            "/api/reconcile",
            json={"root_id": root_id, "ledger_id": ledger_id, "period": None},
        ).json()
        unmarked_row = next(row for row in rerun["rows"] if row["device_no"] == "中-99-99")
        assert unmarked_row["marked_reshot"] is False

        csv_response = client.get(
            "/api/reconcile/export",
            params={"root_id": root_id, "ledger_id": ledger_id, "format": "csv"},
        )
        assert csv_response.status_code == 200
        assert csv_response.content.startswith(b"\xef\xbb\xbf")
        assert csv_response.headers["content-disposition"].endswith('firelens-reshoot.csv"')
        xlsx_response = client.get(
            "/api/reconcile/export",
            params={"root_id": root_id, "ledger_id": ledger_id, "format": "xlsx"},
        )
        assert xlsx_response.status_code == 200
        assert xlsx_response.content.startswith(b"PK")
        assert xlsx_response.headers["content-disposition"].endswith('firelens-reshoot.xlsx"')
