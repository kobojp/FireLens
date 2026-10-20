from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

SECTION_HEADERS = ("棟別樓層", "區段")
DEVICE_HEADERS = ("設備編號", "編號", "滅火器編號", "設備名稱")
SPECIFICATION_HEADERS = ("型式規格",)
LAST_REPLACEMENT_HEADERS = ("上次更換日",)
DEVICE_NO_RE = re.compile(r"^[^\s,]+(?:-[^\s,]+)+$")
REQUIRED_SLOTS = ("編號", "藥劑", "完成")
SPECIAL_PERIODS = {
    "5P 二氧化碳滅火器",
    "壓力錶壞",
    "新增5P 二氧化碳(含放置盒)",
}
DEFAULT_ALIASES = {
    "第一區": "一區",
    "第二區": "二區",
    "B棟異體": "B棟",
}
FLOOR_SUFFIX_RE = re.compile(r"(?:B\d+F|\d+F|1X|1M|RF)$", re.IGNORECASE)
STATUS_LABELS = {
    "complete": "完整",
    "missing_photo": "缺照片",
    "no_folder": "無資料夾",
    "empty_folder": "空資料夾",
    "zero_byte": "0 byte",
    "cloud_not_downloaded": "雲端未下載",
    "invalid_naming": "命名異常",
    "extra_folder": "清冊外資料夾",
    "duplicate_number": "重複編號",
    "near_number": "近似編號",
    "legacy": "舊制／待確認",
    "special": "特例",
}
CLASSIFICATION_LABELS = {
    "complete": "完整",
    "partial_missing": "缺部分照片",
    "no_photos": "完全無照片",
    "no_folder": "找不到設備資料夾",
    "unmapped_legacy": "有未對應／舊制照片",
}
CATEGORY_LABELS = {"extinguisher": "滅火器", "box": "放置盒", "lamp": "燈具"}


class ReconcileError(ValueError):
    pass


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalized_text(value: object) -> str:
    if value is None:
        return ""
    normalized = unicodedata.normalize("NFKC", str(value)).replace("\ufeff", "")
    return re.sub(r"\s+", " ", normalized).strip()


def normalize_device_no(value: object) -> str:
    normalized = _normalized_text(value).upper()
    normalized = re.sub(r"\s*-\s*", "-", normalized)
    return normalized


def normalize_section(value: object) -> str:
    return _normalized_text(value)


def _date_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _text(value)


def _row(
    section: object, device: object, specification: object, replaced: object
) -> dict[str, str]:
    return {
        "section": normalize_section(section),
        "device_no": normalize_device_no(device),
        "specification": _text(specification),
        "last_replacement_date": _date_text(replaced),
    }


def _header_index(values: list[str], aliases: tuple[str, ...]) -> int | None:
    for alias in aliases:
        if alias in values:
            return values.index(alias)
    return None


def _section_device_warning(section: str, device_no: str) -> str | None:
    if not section or "-" not in device_no:
        return None
    prefix = device_no.split("-", 1)[0]
    compact = re.sub(r"\s+", "", unicodedata.normalize("NFKC", section)).upper()
    if prefix and prefix in compact:
        return None
    chinese_ordinals = {"一": "第一", "二": "第二", "三": "第三", "四": "第四"}
    if prefix in chinese_ordinals and compact.startswith(chinese_ordinals[prefix]):
        return None
    return "棟別樓層與設備編號前綴看似不一致，請人工確認"


def _analyze_table(raw_rows: list[list[object]]) -> dict[str, object]:
    material = [row for row in raw_rows if any(_normalized_text(value) for value in row)]
    if not material:
        raise ReconcileError("清冊沒有資料")

    first = [_normalized_text(value) for value in material[0]]
    device_header_index = _header_index(first, DEVICE_HEADERS)
    has_header = device_header_index is not None
    section_header_index = _header_index(first, SECTION_HEADERS) if has_header else None
    specification_header_index = _header_index(first, SPECIFICATION_HEADERS) if has_header else None
    replaced_header_index = _header_index(first, LAST_REPLACEMENT_HEADERS) if has_header else None
    body = material[1:] if has_header else material
    row_offset = 2 if has_header else 1

    parsed_rows: list[dict[str, str]] = []
    preview_rows: list[dict[str, object]] = []
    for offset, values in enumerate(body):
        normalized_values = [_normalized_text(value) for value in values]
        if has_header:
            assert device_header_index is not None
            device_value = values[device_header_index] if device_header_index < len(values) else ""
            section_value = (
                values[section_header_index]
                if section_header_index is not None and section_header_index < len(values)
                else ""
            )
            specification = (
                values[specification_header_index]
                if specification_header_index is not None
                and specification_header_index < len(values)
                else ""
            )
            replaced = (
                values[replaced_header_index]
                if replaced_header_index is not None and replaced_header_index < len(values)
                else ""
            )
        elif len(normalized_values) <= 1:
            section_value = ""
            device_value = values[0] if values else ""
            specification = ""
            replaced = ""
        else:
            section_value = values[0]
            device_value = values[1]
            specification = values[2] if len(values) > 2 else ""
            replaced = values[3] if len(values) > 3 else ""

        row = _row(section_value, device_value, specification, replaced)
        valid = bool(row["device_no"]) and bool(DEVICE_NO_RE.fullmatch(row["device_no"]))
        error = None if valid else "設備編號格式無法辨識"
        warnings: list[str] = []
        if valid and (warning := _section_device_warning(row["section"], row["device_no"])):
            warnings.append(warning)
        preview_rows.append(
            {
                "source_row": row_offset + offset,
                "section": row["section"],
                "device_no": row["device_no"],
                "valid": valid,
                "duplicate": False,
                "warnings": warnings,
                "error": error,
            }
        )
        if valid:
            parsed_rows.append(row)

    counts = Counter(row["device_no"] for row in parsed_rows)
    duplicate_count = 0
    for row in preview_rows:
        if row["valid"] and counts[str(row["device_no"])] > 1:
            row["duplicate"] = True
            duplicate_count += 1
            row["warnings"] = [*row["warnings"], "同一清冊內有重複設備編號"]

    valid_count = sum(bool(row["valid"]) for row in preview_rows)
    invalid_count = len(preview_rows) - valid_count
    if not parsed_rows:
        raise ReconcileError("清冊找不到可辨識的設備編號")
    return {
        "has_header": has_header,
        "total_rows": len(preview_rows),
        "valid_count": valid_count,
        "duplicate_count": duplicate_count,
        "invalid_count": invalid_count,
        "rows": preview_rows,
        "parsed_rows": parsed_rows,
    }


def _public_preview(result: dict[str, object]) -> dict[str, object]:
    return {key: value for key, value in result.items() if key != "parsed_rows"}


def _excel_analysis(data: bytes) -> dict[str, object]:
    if not data:
        raise ReconcileError("Excel 檔案是空的")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ReconcileError("無法讀取 Excel，請確認是有效的 .xlsx 檔案") from exc

    try:
        sheet = workbook.active
        return _analyze_table([list(values) for values in sheet.iter_rows(values_only=True)])
    finally:
        workbook.close()


def parse_excel_bytes(data: bytes) -> list[dict[str, str]]:
    result = _excel_analysis(data)
    return list(result["parsed_rows"])  # type: ignore[arg-type]


def preview_excel_bytes(data: bytes) -> dict[str, object]:
    return _public_preview(_excel_analysis(data))


def _text_analysis(text: str) -> dict[str, object]:
    lines = [line.strip("\ufeff\r") for line in text.splitlines() if line.strip()]
    if not lines:
        raise ReconcileError("貼上的清冊文字是空的")

    delimiter = "\t" if any("\t" in line for line in lines[:3]) else ","
    parsed = list(csv.reader(lines, delimiter=delimiter))
    if len(parsed[0]) == 1 and delimiter == ",":
        parsed = [re.split(r"\s{2,}", line.strip()) for line in lines]
    return _analyze_table([list(values) for values in parsed])


def parse_text_rows(text: str) -> list[dict[str, str]]:
    result = _text_analysis(text)
    return list(result["parsed_rows"])  # type: ignore[arg-type]


def preview_text_rows(text: str) -> dict[str, object]:
    return _public_preview(_text_analysis(text))


def import_ledger(
    conn: sqlite3.Connection,
    *,
    name: str,
    source_type: str,
    rows: list[dict[str, str]],
    source_bytes: bytes | None = None,
    source_text: str | None = None,
    source_filename: str | None = None,
) -> int:
    if source_type not in {"xlsx", "text"}:
        raise ReconcileError("不支援的清冊來源類型")
    if not rows:
        raise ReconcileError("清冊沒有有效資料")
    source_sha256 = hashlib.sha256(source_bytes).hexdigest() if source_bytes is not None else None
    with conn:
        cursor = conn.execute(
            """
            INSERT INTO ledger(name, source_type, source_sha256, source_text, source_filename)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                name.strip() or "未命名清冊",
                source_type,
                source_sha256,
                source_text if source_type == "text" else None,
                source_filename.strip() if source_filename else None,
            ),
        )
        ledger_id = int(cursor.lastrowid)
        conn.executemany(
            """
            INSERT INTO ledger_row(
                ledger_id, row_no, section, device_no, specification,
                last_replacement_date, raw_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    ledger_id,
                    index,
                    row["section"],
                    row["device_no"],
                    row["specification"],
                    row["last_replacement_date"],
                    json.dumps(row, ensure_ascii=False),
                )
                for index, row in enumerate(rows, start=1)
            ],
        )
    return ledger_id


def list_ledgers(conn: sqlite3.Connection) -> list[dict[str, object]]:
    return [
        dict(row)
        for row in conn.execute(
            """
            SELECT l.id, l.name, l.source_type, l.imported_at, l.source_sha256,
                   l.source_filename, COUNT(r.id) AS row_count
            FROM ledger AS l
            LEFT JOIN ledger_row AS r ON r.ledger_id = l.id
            GROUP BY l.id
            ORDER BY l.id DESC
            """
        )
    ]


def get_ledger_detail(conn: sqlite3.Connection, ledger_id: int) -> dict[str, object]:
    row = conn.execute(
        """
        SELECT id, name, source_type, imported_at, source_sha256, source_text, source_filename
        FROM ledger WHERE id=?
        """,
        (ledger_id,),
    ).fetchone()
    if row is None:
        raise ReconcileError("找不到指定清冊")
    result = dict(row)
    result["rows"] = [
        dict(item)
        for item in conn.execute(
            """
            SELECT id, row_no, section, device_no, specification, last_replacement_date, raw_json
            FROM ledger_row WHERE ledger_id=? ORDER BY row_no
            """,
            (ledger_id,),
        )
    ]
    result["row_count"] = len(result["rows"])
    return result


def rename_ledger(conn: sqlite3.Connection, ledger_id: int, name: str) -> dict[str, object]:
    new_name = name.strip()
    if not new_name:
        raise ReconcileError("清冊名稱不可為空")
    if len(new_name) > 200:
        raise ReconcileError("清冊名稱過長")
    current = conn.execute("SELECT name FROM ledger WHERE id=?", (ledger_id,)).fetchone()
    if current is None:
        raise ReconcileError("找不到指定清冊")
    with conn:
        conn.execute("UPDATE ledger SET name=? WHERE id=?", (new_name, ledger_id))
        conn.execute(
            "INSERT INTO audit_log(action, detail_json) VALUES ('rename_ledger', ?)",
            (
                json.dumps(
                    {"ledger_id": ledger_id, "old_name": current["name"], "new_name": new_name},
                    ensure_ascii=False,
                ),
            ),
        )
    return get_ledger_detail(conn, ledger_id)


def delete_ledger(conn: sqlite3.Connection, ledger_id: int) -> dict[str, object]:
    current = conn.execute("SELECT id, name FROM ledger WHERE id=?", (ledger_id,)).fetchone()
    if current is None:
        raise ReconcileError("找不到指定清冊")
    row_count = int(
        conn.execute("SELECT COUNT(*) FROM ledger_row WHERE ledger_id=?", (ledger_id,)).fetchone()[
            0
        ]
    )
    reshoot_count = int(
        conn.execute(
            "SELECT COUNT(*) FROM reshoot_mark WHERE ledger_id=?", (ledger_id,)
        ).fetchone()[0]
    )
    with conn:
        conn.execute("DELETE FROM ledger WHERE id=?", (ledger_id,))
        conn.execute(
            "INSERT INTO audit_log(action, detail_json) VALUES ('delete_ledger', ?)",
            (
                json.dumps(
                    {
                        "ledger_id": ledger_id,
                        "name": current["name"],
                        "row_count": row_count,
                        "reshoot_marks_deleted": reshoot_count,
                    },
                    ensure_ascii=False,
                ),
            ),
        )
    return {
        "deleted": True,
        "ledger_id": ledger_id,
        "name": current["name"],
        "row_count": row_count,
        "reshoot_marks_deleted": reshoot_count,
    }


def seed_default_aliases(conn: sqlite3.Connection) -> None:
    with conn:
        for source, target in DEFAULT_ALIASES.items():
            conn.execute(
                """
                INSERT INTO alias_map(alias_type, source_value, target_value, confirmed)
                VALUES ('building', ?, ?, 1)
                ON CONFLICT(alias_type, source_value) DO NOTHING
                """,
                (source, target),
            )


def list_aliases(conn: sqlite3.Connection) -> list[dict[str, object]]:
    return [
        dict(row)
        for row in conn.execute(
            """
            SELECT id, alias_type, source_value, target_value, confirmed
            FROM alias_map ORDER BY source_value
            """
        )
    ]


def save_alias(
    conn: sqlite3.Connection,
    source_value: str,
    target_value: str,
    *,
    confirmed: bool = True,
) -> dict[str, object]:
    source = source_value.strip()
    target = target_value.strip()
    if not source or not target:
        raise ReconcileError("別名來源與目標都不可為空")
    with conn:
        conn.execute(
            """
            INSERT INTO alias_map(alias_type, source_value, target_value, confirmed)
            VALUES ('building', ?, ?, ?)
            ON CONFLICT(alias_type, source_value)
            DO UPDATE SET target_value = excluded.target_value, confirmed = excluded.confirmed
            """,
            (source, target, int(confirmed)),
        )
    row = conn.execute(
        """
        SELECT id, alias_type, source_value, target_value, confirmed
        FROM alias_map WHERE alias_type='building' AND source_value=?
        """,
        (source,),
    ).fetchone()
    assert row is not None
    return dict(row)


def _alias_map(conn: sqlite3.Connection) -> dict[str, str]:
    return {
        row["source_value"]: row["target_value"]
        for row in conn.execute(
            """
            SELECT source_value, target_value FROM alias_map
            WHERE alias_type='building' AND confirmed=1
            """
        )
    }


def normalize_building(value: str, aliases: dict[str, str]) -> str:
    normalized = unicodedata.normalize("NFKC", value).strip().replace(" ", "")
    for source in sorted(aliases, key=len, reverse=True):
        source_normalized = unicodedata.normalize("NFKC", source).replace(" ", "")
        if normalized.startswith(source_normalized):
            target = unicodedata.normalize("NFKC", aliases[source]).replace(" ", "")
            normalized = target + normalized[len(source_normalized) :]
            break
    normalized = FLOOR_SUFFIX_RE.sub("", normalized)
    return normalized


def _dirty_name_suggestion(filename: str) -> str | None:
    stem = Path(filename).stem.strip()
    cleaned = stem.replace("‵", "").replace("`", "").strip(" ._")
    cleaned = re.sub(r"\s*\(\d+\)$", "", cleaned)
    aliases = {
        "編輯": "編號",
        "編號": "編號",
        "藥劑": "藥劑",
        "有效日期": "藥劑",
        "成": "完成",
        "完成": "完成",
    }
    suggestion = aliases.get(cleaned)
    if suggestion and f"{suggestion}.jpg".casefold() != filename.casefold():
        return f"{suggestion}.jpg"
    return None


def _levenshtein(a: str, b: str) -> int:
    if a == b:
        return 0
    if not a:
        return len(b)
    previous = list(range(len(b) + 1))
    for i, char_a in enumerate(a, start=1):
        current = [i]
        for j, char_b in enumerate(b, start=1):
            current.append(
                min(
                    current[-1] + 1,
                    previous[j] + 1,
                    previous[j - 1] + (char_a != char_b),
                )
            )
        previous = current
    return previous[-1]


def _near_candidates(device_no: str, item_codes: list[str]) -> list[str]:
    target_parts = device_no.split("-")
    if len(target_parts) < 2:
        return []
    candidates: list[str] = []
    for code in item_codes:
        parts = code.split("-")
        if len(parts) != len(target_parts):
            continue
        if parts[:-1] != target_parts[:-1]:
            continue
        target_tail = target_parts[-1]
        candidate_tail = parts[-1]
        if abs(len(candidate_tail) - len(target_tail)) != 1:
            continue
        if _levenshtein(candidate_tail, target_tail) == 1:
            candidates.append(code)
    return sorted(set(candidates))[:5]


def _item_payload(conn: sqlite3.Connection, item: sqlite3.Row) -> dict[str, object]:
    photos = list(
        conn.execute(
            """
            SELECT id, filename, relative_path, slot, size_bytes, cloud_placeholder
            FROM photo WHERE item_id=? ORDER BY filename COLLATE NOCASE
            """,
            (item["id"],),
        )
    )
    slots = {row["slot"] for row in photos if row["slot"]}
    unmapped_count = sum(row["slot"] is None for row in photos)
    missing_slots = [slot for slot in REQUIRED_SLOTS if slot not in slots]
    suggestions = [
        {"filename": row["filename"], "suggestion": suggestion}
        for row in photos
        if row["slot"] is None and (suggestion := _dirty_name_suggestion(row["filename"]))
    ]
    if any(row["cloud_placeholder"] for row in photos):
        status = "cloud_not_downloaded"
    elif any(int(row["size_bytes"]) <= 0 for row in photos):
        status = "zero_byte"
    elif suggestions:
        status = "invalid_naming"
    elif missing_slots:
        status = "missing_photo"
    else:
        status = "complete"
    if not photos:
        status = "empty_folder"
    return {
        "item_id": item["id"],
        "year": item["year"],
        "period": item["period"],
        "building": item["building"],
        "relative_path": item["relative_path"],
        "photo_count": item["photo_count"],
        "missing_slots": missing_slots,
        "naming_suggestions": suggestions,
        "unmapped_count": unmapped_count,
        "status": status,
    }


def _slot_only_payload(conn: sqlite3.Connection, item: sqlite3.Row) -> dict[str, object]:
    photos = list(
        conn.execute(
            """
            SELECT id, filename, relative_path, slot, size_bytes, cloud_placeholder
            FROM photo WHERE item_id=? ORDER BY filename COLLATE NOCASE
            """,
            (item["id"],),
        )
    )
    required_slots = tuple(json.loads(item["slots_json"]))
    unmapped_count = sum(row["slot"] is None for row in photos)
    present_slots = {row["slot"] for row in photos if row["slot"]}
    missing_slots = [slot for slot in required_slots if slot not in present_slots]
    if item["period"] in SPECIAL_PERIODS:
        status = "special"
        missing_slots = []
    elif item["status"] == "legacy" or (
        item["category_code"] == "lamp" and item["status"] == "unknown"
    ):
        status = "legacy"
        missing_slots = []
    elif not photos:
        status = "empty_folder"
    elif any(row["cloud_placeholder"] for row in photos):
        status = "cloud_not_downloaded"
    elif any(int(row["size_bytes"]) <= 0 for row in photos):
        status = "zero_byte"
    elif not missing_slots:
        status = "complete"
    elif present_slots:
        status = "missing_photo"
    else:
        status = "invalid_naming"
    return {
        "item_id": item["id"],
        "year": item["year"],
        "period": item["period"],
        "building": item["building"],
        "relative_path": item["relative_path"],
        "photo_count": item["photo_count"],
        "missing_slots": missing_slots,
        "naming_suggestions": [],
        "unmapped_count": unmapped_count,
        "status": status,
    }


def _result_classification(row: dict[str, object]) -> str:
    status = str(row.get("status", ""))
    found = row.get("found", [])
    has_unmapped = any(int(item.get("unmapped_count", 0)) > 0 for item in found)  # type: ignore[union-attr]
    if status in {"complete", "special"} and not has_unmapped:
        return "complete"
    if status in {"legacy", "invalid_naming"} or has_unmapped:
        return "unmapped_legacy"
    if status == "empty_folder":
        return "no_photos"
    if status in {"no_folder", "near_number", "extra_folder"}:
        return "no_folder"
    return "partial_missing"


def set_reshoot_mark(
    conn: sqlite3.Connection,
    *,
    root_id: int,
    ledger_id: int,
    category_code: str,
    device_no: str,
    marked: bool,
) -> None:
    if not device_no.strip():
        raise ReconcileError("設備編號不可為空")
    with conn:
        conn.execute(
            """
            INSERT INTO reshoot_mark(
                root_id, ledger_id, category_code, device_no, marked, marked_at
            ) VALUES (?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(root_id, ledger_id, category_code, device_no)
            DO UPDATE SET marked=excluded.marked, marked_at=CURRENT_TIMESTAMP
            """,
            (root_id, ledger_id, category_code, device_no, int(marked)),
        )


def _reshoot_marks(conn: sqlite3.Connection, root_id: int, ledger_id: int) -> set[tuple[str, str]]:
    return {
        (row["category_code"], row["device_no"])
        for row in conn.execute(
            """
            SELECT category_code, device_no FROM reshoot_mark
            WHERE root_id=? AND ledger_id=? AND marked=1
            """,
            (root_id, ledger_id),
        )
    }


def build_reconciliation(
    conn: sqlite3.Connection,
    *,
    root_id: int,
    ledger_id: int,
    period: str | None = None,
) -> dict[str, object]:
    seed_default_aliases(conn)
    root = conn.execute("SELECT id FROM root WHERE id=?", (root_id,)).fetchone()
    if root is None:
        raise ReconcileError("找不到指定照片根目錄")
    ledger = conn.execute("SELECT id, name FROM ledger WHERE id=?", (ledger_id,)).fetchone()
    if ledger is None:
        raise ReconcileError("找不到指定清冊")

    aliases = _alias_map(conn)
    marked_pairs = _reshoot_marks(conn, root_id, ledger_id)
    ledger_rows = list(
        conn.execute(
            """
            SELECT row_no, section, device_no, specification, last_replacement_date
            FROM ledger_row WHERE ledger_id=? ORDER BY row_no
            """,
            (ledger_id,),
        )
    )
    item_query = """
        SELECT i.id, i.code, i.year, i.period, i.building, i.relative_path,
               i.status, i.photo_count
        FROM item AS i
        JOIN category AS c ON c.id=i.category_id
        WHERE i.root_id=? AND c.code='extinguisher'
    """
    params: list[object] = [root_id]
    if period:
        item_query += " AND i.period=?"
        params.append(period)
    items = list(conn.execute(item_query, params))

    items_by_code: dict[str, list[sqlite3.Row]] = defaultdict(list)
    for item in items:
        items_by_code[item["code"]].append(item)
    item_codes = list(items_by_code)
    ledger_ids = [row["device_no"] for row in ledger_rows]
    ledger_id_counts = Counter(ledger_ids)
    rows: list[dict[str, object]] = []

    for ledger_row in ledger_rows:
        device_no = ledger_row["device_no"]
        expected_building = normalize_building(ledger_row["section"] or "", aliases)
        matches = items_by_code.get(device_no, [])
        base: dict[str, object] = {
            "kind": "ledger",
            "category": "extinguisher",
            "row_no": ledger_row["row_no"],
            "section": ledger_row["section"] or "",
            "expected_building": expected_building,
            "device_no": device_no,
            "specification": ledger_row["specification"] or "",
            "last_replacement_date": ledger_row["last_replacement_date"] or "",
            "found": [],
            "missing_slots": [],
            "naming_suggestions": [],
            "near_candidates": [],
            "marked_reshot": ("extinguisher", device_no) in marked_pairs,
        }
        if ledger_id_counts[device_no] > 1 or len(matches) > 1:
            base["status"] = "duplicate_number"
            base["found"] = [_item_payload(conn, item) for item in matches]
        elif not matches:
            near = _near_candidates(device_no, item_codes)
            base["near_candidates"] = near
            base["status"] = "near_number" if near else "no_folder"
        else:
            payload = _item_payload(conn, matches[0])
            base["status"] = payload["status"]
            base["found"] = [payload]
            base["missing_slots"] = payload["missing_slots"]
            base["naming_suggestions"] = payload["naming_suggestions"]
            actual_building = normalize_building(matches[0]["building"], aliases)
            base["building_match"] = not expected_building or expected_building == actual_building
        base["classification"] = _result_classification(base)
        rows.append(base)

    ledger_set = set(ledger_ids)
    for code, code_items in sorted(items_by_code.items()):
        if code in ledger_set:
            continue
        for item in code_items:
            payload = _item_payload(conn, item)
            is_special = item["period"] in SPECIAL_PERIODS
            extra_row: dict[str, object] = {
                "kind": "special" if is_special else "extra",
                "category": "extinguisher",
                "row_no": None,
                "section": "",
                "expected_building": "",
                "device_no": code,
                "specification": "",
                "last_replacement_date": "",
                "status": "special" if is_special else "extra_folder",
                "found": [payload],
                "missing_slots": [] if is_special else payload["missing_slots"],
                "naming_suggestions": ([] if is_special else payload["naming_suggestions"]),
                "near_candidates": [],
                "marked_reshot": ("extinguisher", code) in marked_pairs,
            }
            extra_row["classification"] = _result_classification(extra_row)
            rows.append(extra_row)

    slot_query = """
        SELECT i.id, i.code, i.year, i.period, i.building, i.relative_path,
               i.status, i.photo_count, c.code AS category_code, c.name AS category_name,
               c.slots_json
        FROM item AS i
        JOIN category AS c ON c.id=i.category_id
        WHERE i.root_id=? AND c.code IN ('box', 'lamp')
    """
    slot_params: list[object] = [root_id]
    if period:
        slot_query += " AND i.period=?"
        slot_params.append(period)
    for slot_item in conn.execute(slot_query, slot_params):
        payload = _slot_only_payload(conn, slot_item)
        slot_row: dict[str, object] = {
            "kind": "slot_only",
            "category": slot_item["category_code"],
            "row_no": None,
            "section": slot_item["building"],
            "expected_building": normalize_building(slot_item["building"], aliases),
            "device_no": slot_item["code"],
            "specification": slot_item["category_name"],
            "last_replacement_date": "",
            "status": payload["status"],
            "found": [payload],
            "missing_slots": payload["missing_slots"],
            "naming_suggestions": [],
            "near_candidates": [],
            "marked_reshot": (slot_item["category_code"], slot_item["code"]) in marked_pairs,
        }
        slot_row["classification"] = _result_classification(slot_row)
        rows.append(slot_row)

    ledger_result_rows = [row for row in rows if row["kind"] == "ledger"]
    counts = Counter(str(row["status"]) for row in rows)
    classification_counts = Counter(str(row["classification"]) for row in ledger_result_rows)
    complete_count = sum(row["classification"] == "complete" for row in ledger_result_rows)
    total = len(ledger_result_rows)
    progress_acc: dict[str, Counter[str]] = defaultdict(Counter)
    for row in ledger_result_rows:
        building = str(row["section"] or "未提供棟別樓層")
        progress_acc[building]["total"] += 1
        if row["classification"] == "complete":
            progress_acc[building]["complete"] += 1
    building_progress = [
        {
            "building": building,
            "total": values["total"],
            "complete": values["complete"],
            "rate": round(values["complete"] * 100 / values["total"], 1)
            if values["total"]
            else 0.0,
        }
        for building, values in sorted(progress_acc.items())
    ]
    return {
        "root_id": root_id,
        "ledger_id": ledger_id,
        "ledger_name": ledger["name"],
        "period": period,
        "total": total,
        "complete": complete_count,
        "completion_rate": round(complete_count * 100 / total, 1) if total else 0.0,
        "counts": dict(sorted(counts.items())),
        "classification_counts": dict(sorted(classification_counts.items())),
        "building_progress": building_progress,
        "rows": rows,
    }


def reshoot_rows(result: dict[str, object]) -> list[dict[str, object]]:
    return [
        row
        for row in result["rows"]  # type: ignore[index]
        if row.get("classification") != "complete" and row.get("kind") != "special"
    ]


def reshoot_text(result: dict[str, object]) -> str:
    lines: list[str] = []
    for row in reshoot_rows(result):
        missing = "、".join(row.get("missing_slots", []))
        classification_label = CLASSIFICATION_LABELS.get(
            str(row.get("classification", "")), str(row.get("classification", ""))
        )
        detail = f"缺：{missing}" if missing else classification_label
        category = CATEGORY_LABELS.get(str(row.get("category", "")), str(row.get("category", "")))
        lines.append(f"{category}\t{row['device_no']}\t{row.get('section', '')}\t{detail}")
    return "\n".join(lines)


def export_reconciliation(result: dict[str, object], file_format: str) -> tuple[bytes, str, str]:
    rows = reshoot_rows(result)
    headers = [
        "類別",
        "設備編號",
        "棟別樓層",
        "核對分類",
        "缺少照片／狀態",
        "找到月份",
        "找到棟別",
        "相對路徑",
        "近似編號",
        "已補拍",
    ]
    body_rows: list[list[str]] = []
    for row in rows:
        found = row.get("found", [])
        found_periods = "、".join(str(item.get("period", "")) for item in found)
        found_buildings = "、".join(str(item.get("building", "")) for item in found)
        paths = "、".join(str(item.get("relative_path", "")) for item in found)
        body_rows.append(
            [
                CATEGORY_LABELS.get(str(row.get("category", "")), str(row.get("category", ""))),
                str(row.get("device_no", "")),
                str(row.get("section", "")),
                CLASSIFICATION_LABELS.get(
                    str(row.get("classification", "")), str(row.get("classification", ""))
                ),
                "、".join(row.get("missing_slots", []))
                or STATUS_LABELS.get(str(row.get("status", "")), str(row.get("status", ""))),
                found_periods,
                found_buildings,
                paths,
                "、".join(row.get("near_candidates", [])),
                "是" if row.get("marked_reshot") else "否",
            ]
        )

    if file_format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(body_rows)
        return stream.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", "補拍清單.csv"
    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "補拍清單"
        sheet.append(headers)
        for body in body_rows:
            sheet.append(body)
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
            sheet.column_dimensions[column[0].column_letter].width = width
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "補拍清單.xlsx",
        )
    raise ReconcileError("匯出格式只支援 csv 或 xlsx")
